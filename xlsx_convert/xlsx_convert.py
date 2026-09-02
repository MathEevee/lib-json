import csv
import json
import os
import datetime
from openpyxl import Workbook, load_workbook


def _dedupe_headers(raw_headers):
    """
    Make headers unique.

    Empty/None header cells become "col", "col_1", "col_2", ... and
    any duplicated header name gets a numeric suffix so no column is
    lost when the row is turned into a dict. Non-string header values
    (e.g. a percentage stored as a raw number) are stringified first.

    Args:
        raw_headers (list | tuple): The raw header row read from the
            CSV file or worksheet.

    Returns:
        list: Deduplicated, non-empty header names, same length/order
            as raw_headers.
    """
    headers = []
    seen = {}
    for h in raw_headers:
        name = str(h) if h not in (None, "") else "col"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        headers.append(name)
    return headers


def _fieldnames_union(rows):
    """
    Collect every key found across a list of dicts, in first-seen order.

    Used so the output sheet/CSV has a column for every field that
    appears in ANY row, even if some rows don't have that field.

    Args:
        rows (list): List of dictionaries.

    Returns:
        list: Ordered list of unique field names.
    """
    fieldnames = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                fieldnames.append(key)
    return fieldnames


def _write_sheet(ws, rows):
    """
    Write a list of dicts into a worksheet as a header row + data rows.

    Args:
        ws (Worksheet): The openpyxl worksheet to write into.
        rows (list): List of dictionaries representing the rows.
    """
    if not rows:
        return
    fieldnames = _fieldnames_union(rows)
    ws.append(fieldnames)
    for row in rows:
        ws.append([row.get(field, "") for field in fieldnames])


def _make_json_safe(value):
    """
    Convert a cell value into something json.dump can serialize.

    openpyxl can return datetime.date/datetime.datetime/datetime.time
    objects for date-formatted cells; these get converted to ISO 8601
    strings.

    Args:
        value: The raw cell value.

    Returns:
        A JSON-serializable value.
    """
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return value


def _sheet_to_rows(ws):
    """
    Turn a worksheet into a list of row dicts, keyed by deduped headers.

    Args:
        ws (Worksheet): The openpyxl worksheet to read (read-only mode).

    Returns:
        list: One dict per data row (header row excluded).
    """
    rows_iter = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration:
        return []
    headers = _dedupe_headers(raw_headers)
    return [
        {h: _make_json_safe(v) for h, v in zip(headers, row)}
        for row in rows_iter
    ]


def to_xlsx(input_file, output_xlsx=None):
    """
    Convert a CSV or JSON file into an XLSX workbook.

    - If the source is a CSV file, the output has a single sheet.
    - If the source is a JSON file containing a list of objects, the
      output has a single sheet.
    - If the source is a JSON file containing a dict of lists
      (e.g. {"Sheet1": [...], "Sheet2": [...]}), each key becomes its
      own sheet in the workbook.

    Args:
        input_file (str): Path to the source .csv or .json file.
        output_xlsx (str, optional): Path to the output XLSX file.
            Defaults to the input filename with a .xlsx extension.

    Returns:
        str: Path to the generated XLSX file.

    Raises:
        ValueError: If the file extension or JSON structure is not
            supported, or if there is no data to convert.
    """
    if output_xlsx is None:
        output_xlsx = os.path.splitext(input_file)[0] + ".xlsx"

    ext = os.path.splitext(input_file)[1].lower()
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    if ext == ".csv":
        with open(input_file, mode="r", newline="", encoding="utf-8") as csvfile:
            reader = csv.reader(csvfile)
            raw_headers = next(reader)
            headers = _dedupe_headers(raw_headers)
            data = [dict(zip(headers, row)) for row in reader]

        sheet_title = os.path.splitext(os.path.basename(input_file))[0][:31]
        ws = wb.create_sheet(title=sheet_title)
        _write_sheet(ws, data)

    elif ext == ".json":
        with open(input_file, "r", encoding="utf-8") as jsonfile:
            data = json.load(jsonfile)

        if isinstance(data, dict):
            # One sheet per key
            for sheet_name, rows in data.items():
                ws = wb.create_sheet(title=str(sheet_name)[:31])  # Excel caps titles at 31 chars
                _write_sheet(ws, rows)
        elif isinstance(data, list):
            ws = wb.create_sheet(title="Sheet1")
            _write_sheet(ws, data)
        else:
            raise ValueError("Unsupported JSON structure for XLSX conversion.")

    else:
        raise ValueError(f"Unsupported extension '{ext}' for XLSX conversion.")

    if not wb.sheetnames:
        raise ValueError("No data to convert.")

    wb.save(output_xlsx)
    return output_xlsx


def to_json(input_xlsx, output_json=None):
    """
    Convert an XLSX workbook into one JSON file per sheet.

    Each sheet is written to its own file, named
    "<input_basename>_p<sheet_number>.json" (e.g. "skycreep_p1.json",
    "skycreep_p2.json", ...). If output_json is provided, it is used
    as the base path instead of deriving one from input_xlsx.

    Args:
        input_xlsx (str): Path to the source XLSX file.
        output_json (str, optional): Base path to derive per-sheet
            output filenames from. Defaults to input_xlsx's path
            with its extension stripped.

    Returns:
        list: Paths to all the generated JSON files.
    """
    base_path = (
        os.path.splitext(output_json)[0]
        if output_json
        else os.path.splitext(input_xlsx)[0]
    )

    wb = load_workbook(filename=input_xlsx, read_only=True, data_only=True)

    output_paths = []
    for i, sheet_name in enumerate(wb.sheetnames, start=1):
        data = _sheet_to_rows(wb[sheet_name])
        output_path = f"{base_path}_p{i}.json"
        with open(output_path, "w", encoding="utf-8") as jsonfile:
            json.dump(data, jsonfile, ensure_ascii=False, indent=4)
        output_paths.append(output_path)

    return output_paths